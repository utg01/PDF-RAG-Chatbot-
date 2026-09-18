import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment


BASELINE_PATH = "evaluation/results/baseline_ragas_results.json"
ADVANCED_PATH = "evaluation/results/advanced_ragas_results.json"

OUTPUT_PATH = "evaluation/results/ragas_comparison.xlsx"


def load_json(path):
    if not os.path.exists(path):
        print(f"File not found: {path}")
        return []

    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def format_results(data):

    if not data:
        return "Not evaluated"

    # RAGAS scores are nested inside "scores"
    scores = data.get("scores", {})

    def fmt(value):
        try:
            if value is None:
                return "N/A"

            return f"{float(value):.4f}"

        except (ValueError, TypeError):
            return "N/A"

    return (
        f"Faithfulness: {fmt(scores.get('faithfulness'))}\n"
        f"Answer Relevancy: {fmt(scores.get('answer_relevancy'))}\n"
        f"Context Precision: {fmt(scores.get('context_precision'))}\n"
        f"Context Recall: {fmt(scores.get('context_recall'))}"
    )


def main():

    # Load JSON files
    baseline = load_json(BASELINE_PATH)
    advanced = load_json(ADVANCED_PATH)

    # Create maps using question ID
    baseline_map = {
        item["id"]: item
        for item in baseline
        if "id" in item
    }

    advanced_map = {
        item["id"]: item
        for item in advanced
        if "id" in item
    }

    # Get all unique question IDs
    all_ids = sorted(
        set(baseline_map) |
        set(advanced_map)
    )

    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "RAGAS Comparison"

    # Headers
    ws.append([
        "Question ID",
        "Question",
        "Baseline Results",
        "Advanced Results"
    ])

    # Header formatting
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # Add question results
    for question_id in all_ids:

        baseline_result = baseline_map.get(question_id)
        advanced_result = advanced_map.get(question_id)

        # Get question text from whichever result exists
        question = ""

        if baseline_result:
            question = baseline_result.get("question", "")

        elif advanced_result:
            question = advanced_result.get("question", "")

        ws.append([
            question_id,
            question,
            format_results(baseline_result),
            format_results(advanced_result)
        ])

    # Freeze header
    ws.freeze_panes = "A2"

    # Column widths
    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 65
    ws.column_dimensions["C"].width = 35
    ws.column_dimensions["D"].width = 35

    # Format cells
    for row in ws.iter_rows(min_row=2):

        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    # Set row heights
    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 80

    # Create output directory
    output_dir = os.path.dirname(OUTPUT_PATH)

    if output_dir:
        os.makedirs(
            output_dir,
            exist_ok=True
        )

    # Save workbook
    wb.save(OUTPUT_PATH)

    # Console output
    print("\n========================================")
    print("RAGAS Comparison Excel Created")
    print("========================================")
    print(f"Output: {OUTPUT_PATH}")
    print(f"Baseline results: {len(baseline)}")
    print(f"Advanced results: {len(advanced)}")
    print(f"Total unique questions: {len(all_ids)}")
    print("========================================\n")


if __name__ == "__main__":
    main()